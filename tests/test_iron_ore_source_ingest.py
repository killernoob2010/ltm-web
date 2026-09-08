from datetime import datetime
from pathlib import Path

import openpyxl


def _inventory_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "总览"
    ws.append(["统计日期", "港口", "区域", "库存总量", "粉矿", "块矿", "球团", "精粉"])
    ws.append([datetime(2026, 9, 1), "样本2", "沿江", 8, 3, 1, 2, 2])
    ws.append([datetime(2026, 9, 1), "总计", "", 100, 60, 15, 10, 15])

    grade = wb.create_sheet("分品位")
    grade.append([None, None, None, "Fe 64%以上", None, None, None, None, "Fe60%-64%", None, None, None, None, "Fe 55%-60%", None, None, None, None, "Fe 55%以下", None])
    grade.append(["统计日期", "港口", "区域", "高品粉矿", "高品块矿", "高品球团", "高品精粉", "高品总计", "中高品粉矿", "中高品块矿", "中高品球团", "中高品精粉", "中高品总计", "中低品粉矿", "中低品块矿", "中低品球团", "中低品精粉", "中低品总计", "低品粉矿", "低品块矿"])
    grade.append([datetime(2026, 9, 1), "样本2", "沿江", 1, 0, 0, 0, 1, 2, 0, 0, 0, 2, 0, 0, 0, 0, 0, 3, 0])

    powder = wb.create_sheet("粗粉")
    powder.append(["统计日期", "港口", "区域", "粉矿总计", "澳粉总计", "PB粉", "纽曼粉"])
    powder.append([datetime(2026, 9, 1), "样本2", "沿江", 3, 3, 1, None])
    powder.append([datetime(2026, 9, 1), "总计", "", 60, 60, 20, 10])
    for sheet_name, total_header in (("块矿", "块矿总计"), ("球团", "球团总计"), ("精粉", "精粉总计")):
        sheet = wb.create_sheet(sheet_name)
        sheet.append(["统计日期", "港口", "区域", total_header, "PB块" if sheet_name == "块矿" else "澳大利亚"])
        sheet.append([datetime(2026, 9, 1), "样本2", "沿江", 1, 1])
    wb.save(path)


def _arrival_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "国家"
    ws.append(["全国到港量来源地货量明细"])
    ws.append([])
    ws.append(["到港时间", "港口", "总计", "澳大利亚", "巴西"])
    ws.append([datetime(2026, 8, 30), "南通港", 100, 70, 30])
    ws.append([datetime(2026, 8, 30), "日照港", 200, 120, None])

    product = wb.create_sheet("品种")
    product.append(["全国到港量分货物明细"])
    product.append([])
    product.append(["到港时间", "港口", "总计", "PB粉", "卡粉"])
    product.append([datetime(2026, 8, 30), "南通港", 100, 40, 60])

    form = wb.create_sheet("货种品位")
    form.append(["全国到港量分货种及品位货物明细"])
    form.append([])
    form.append([None, None, None, "粉矿"])
    form.append(["到港时间", "港口", "总计", "60%以下", "粉矿 汇总"])
    form.append([datetime(2026, 8, 30), "南通港", 100, None, 100])
    wb.save(path)


def test_parse_source_files_captures_port_product_grade_and_actual_arrival(tmp_path):
    from backend.app.iron_ore_source_ingest import parse_mysteel_source_files

    inventory = tmp_path / "库存.xlsx"
    arrival = tmp_path / "到港.xlsx"
    _inventory_workbook(inventory)
    _arrival_workbook(arrival)

    package = parse_mysteel_source_files([inventory, arrival])

    nantong = [row for row in package.inventory_port_product if row["port_name"] == "南通" and row["product"] == "PB粉"]
    assert nantong and nantong[0]["sample_name"] == "样本2"
    assert nantong[0]["scope_type"] == "sample"
    assert nantong[0]["value"] == 1.0
    assert any(row["grade"] == "高品" and row["value"] == 1.0 for row in package.inventory_grade)

    actual = [row for row in package.arrival_actual if row["slice_type"] == "product" and row["product"] == "PB粉"]
    assert actual and actual[0]["arrival_kind"] == "actual"
    assert actual[0]["port_name"] == "南通"

    missing_country = [row for row in package.arrival_actual if row["slice_type"] == "country" and row["dimension"] == "巴西" and row["port_name"] == "日照"]
    assert missing_country and missing_country[0]["value"] is None
    assert missing_country[0]["value_status"] == "missing"


def test_build_v2_workbook_round_trips_authoritative_sheets(tmp_path):
    from backend.app.iron_ore_source_ingest import build_v2_workbook, parse_mysteel_source_files

    inventory = tmp_path / "库存.xlsx"
    arrival = tmp_path / "到港.xlsx"
    _inventory_workbook(inventory)
    _arrival_workbook(arrival)
    package = parse_mysteel_source_files([inventory, arrival])

    output = tmp_path / "整合V2.xlsx"
    output.write_bytes(build_v2_workbook(package))
    wb = openpyxl.load_workbook(output, read_only=True, data_only=True)
    assert {"整合明细", "港口品种库存", "库存汇总分档", "到港明细", "批次信息"}.issubset(wb.sheetnames)
    assert any(row[3] == "南通" and row[7] == "PB粉" for row in wb["港口品种库存"].iter_rows(min_row=2, values_only=True))
    assert any(row[0] == "actual" for row in wb["到港明细"].iter_rows(min_row=2, values_only=True))
    wb.close()


def test_historical_inventory_first_block_and_old_grades(tmp_path):
    from backend.app.iron_ore_source_ingest import parse_mysteel_source_files
    path = tmp_path / 'history.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '总览'
    ws.append(['统计日期', '港口', '区域', '库存总量'])
    ws.append([datetime(2025, 9, 1), 'Total', '', 10])
    ws = wb.create_sheet('粉矿')
    ws.append(['统计日期', '港口', '区域', 'PB粉', '统计日期', '港口', '区域', 'PB粉'])
    ws.append([datetime(2025, 9, 1), 'Total', '', 10, datetime(2025, 9, 1), 'Total', '', 999])
    ws = wb.create_sheet('分品位')
    ws.append(['统计日期', '港口', '区域', '高品总计', '中品总计', '低品总计'])
    ws.append([datetime(2025, 9, 1), 'Jiangyin', '', 5, 3, 2])
    wb.save(path)
    p = parse_mysteel_source_files([path])
    assert len(p.inventory_port_product) == 1
    assert p.inventory_port_product[0]['value'] == 10
    assert p.inventory_port_product[0]['scope_type'] == 'total'
    assert p.inventory_port_product[0]['category'] == '粉矿'
    assert {r['grade'] for r in p.inventory_grade} == {'高品（旧三档）', '中低品', '低品'}
    assert {r['port_name'] for r in p.inventory_grade} == {'江阴'}


def test_arrival_repeated_grade_labels_keep_parent_cargo_form(tmp_path):
    from backend.app.iron_ore_source_ingest import parse_mysteel_source_files
    path = tmp_path / 'arrival.xlsx'
    _arrival_workbook(path)
    wb = openpyxl.load_workbook(path)
    ws = wb['货种品位']
    ws.delete_rows(1, ws.max_row)
    ws.append(['全国到港量分货种及品位货物明细'])
    ws.append([])
    ws.append([None, None, None, '粉矿', None, '块矿', None])
    ws.append(['到港时间', '港口', '总计', '60%以下', '粉矿 汇总', '60%以下', '块矿 汇总'])
    ws.append([datetime(2026, 8, 30), '南通港', 100, 30, 60, 10, 40])
    wb.save(path)
    rows = [r for r in parse_mysteel_source_files([path]).arrival_actual if r['slice_type']=='form' and r['dimension']=='60%以下']
    assert {(r['category'], r['value']) for r in rows} == {('粉矿', 30), ('块矿', 10)}


def test_estimated_arrival_preserves_distinct_source_columns_for_same_product(tmp_path):
    from backend.app.iron_ore_source_ingest import parse_mysteel_source_files
    path = tmp_path / 'forecast.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '澳洲预计到达中国锚地量'
    ws.append(['日期', '杨迪粉', '大杨迪', 'PB粉'])
    ws.append(['2026/08/24-2026/08/30', 1, 2, 3])
    wb.save(path)
    rows = [r for r in parse_mysteel_source_files([path]).arrival_estimated if r['product']=='杨迪粉']
    assert {(r['dimension'], r['value'], r['source_cell']) for r in rows} == {('杨迪粉', 1, 'B2'), ('大杨迪', 2, 'C2')}


def test_historical_inventory_withholds_unreliable_totals_without_changing_raw_values():
    from backend.app.iron_ore_source_ingest import SourcePackage, _flag_historical_inventory
    p = SourcePackage()
    base = {'source_file': 'history.xlsx', 'observed_date': '2020-01-07', 'category': '粉矿', 'source_country': '澳洲', 'value_status': 'numeric'}
    for product, total in [('PB粉', 15), ('麦克粉', 30), ('卡拉拉精粉', 15)]:
        p.inventory_port_product.append(dict(base, product=product, scope_type='total', value=total))
        p.inventory_port_product.extend(dict(base, product=product, scope_type='sample', port_name=str(i), value=1) for i in range(15))
    _flag_historical_inventory(p, 'history.xlsx')
    totals = [r for r in p.inventory_port_product if r['scope_type'] == 'total']
    assert [(r['value'], r['value_status']) for r in totals] == [(15, 'numeric'), (30, 'withheld_total_mismatch'), (15, 'withheld_historical_column')]
    assert all(r['value'] == 1 for r in p.inventory_port_product if r['scope_type'] == 'sample')


def test_v2_roundtrip_preserves_original_source_row(tmp_path):
    from backend.app.iron_ore_source_ingest import SourcePackage, build_v2_workbook
    from backend.app.data_visualization import _parse_integrated_excel
    p = SourcePackage()
    p.inventory_port_product = [{'observed_date': '2020-01-07', 'week_start': '2020-01-06', 'scope_type': 'total', 'product': 'PB粉', 'category': '粉矿', 'value': 15, 'value_status': 'numeric', 'source_file': 'history.xlsx', 'source_sheet': '粉矿', 'source_row': 4567, 'source_column': 4, 'source_cell': 'D4567'}]
    path = tmp_path / 'history-v2.xlsx'
    path.write_bytes(build_v2_workbook(p))
    result = _parse_integrated_excel(path)
    assert not result['errors']
    row = result['details']['inventory_port_product'][0]
    assert (row['source_row'], row['source_column'], row['source_cell']) == (4567, 4, 'D4567')

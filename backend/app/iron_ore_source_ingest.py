"""Source-file parsing for the iron-ore weekly report V2 workflow.

The existing data-visualization tables intentionally keep their V1 summary
shape.  This module reads the original Mysteel workbooks before they are
flattened, preserving the port/sample, product, grade, source-cell and value
status needed by the report and by later reconciliation work.
"""

from __future__ import annotations

import hashlib
import io
import json
import shutil
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

from . import db


PARSER_VERSION = "iron-ore-source-v2.1"
MAPPING_VERSION = "mysteel-port-sample-2026-09"
STRUCTURE_VERSION = "iron-ore-integrated-v2"

# The inventory workbook uses sample labels instead of port names.  This is a
# controlled mapping for the 15-port inventory universe; the raw sample label
# is always retained on each row so a future mapping revision is auditable.
SAMPLE_TO_PORT = {
    "样本1": ("江阴", "沿江"),
    "样本2": ("南通", "沿江"),
    "样本3": ("太仓", "沿江"),
    "样本4": ("舟山", "沿江"),
    "样本5": ("福州", "华南"),
    "样本6": ("湛江", "华南"),
    "样本7": ("岚桥", "沿海/北方"),
    "样本8": ("岚山", "沿海/北方"),
    "样本9": ("连云港", "沿海/北方"),
    "样本10": ("青岛", "沿海/北方"),
    "样本11": ("日照", "沿海/北方"),
    "样本12": ("曹妃甸", "沿海/北方"),
    "样本13": ("黄骅", "沿海/北方"),
    "样本14": ("京唐", "沿海/北方"),
    "样本15": ("天津", "沿海/北方"),
}


@dataclass
class SourcePackage:
    """Parsed, not-yet-activated source data."""

    package_id: str = field(default_factory=lambda: uuid.uuid4().hex)
    parser_version: str = PARSER_VERSION
    mapping_version: str = MAPPING_VERSION
    structure_version: str = STRUCTURE_VERSION
    source_files: List[Dict[str, Any]] = field(default_factory=list)
    inventory_port_product: List[Dict[str, Any]] = field(default_factory=list)
    inventory_summary: List[Dict[str, Any]] = field(default_factory=list)
    inventory_grade: List[Dict[str, Any]] = field(default_factory=list)
    inventory_mainstream: List[Dict[str, Any]] = field(default_factory=list)
    arrival_actual: List[Dict[str, Any]] = field(default_factory=list)
    arrival_estimated: List[Dict[str, Any]] = field(default_factory=list)
    shipments: List[Dict[str, Any]] = field(default_factory=list)
    legacy_points: List[Dict[str, Any]] = field(default_factory=list)
    validation: Dict[str, Any] = field(default_factory=dict)

    @property
    def all_rows(self) -> Iterable[Dict[str, Any]]:
        yield from self.inventory_port_product
        yield from self.inventory_summary
        yield from self.inventory_grade
        yield from self.inventory_mainstream
        yield from self.arrival_actual
        yield from self.arrival_estimated
        yield from self.shipments


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _date_value(value: Any) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    text = _text(value)
    for candidate in (text[:10], text.split("-")[0].strip()):
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(candidate, fmt).date()
            except ValueError:
                continue
    return None


def _number(value: Any) -> tuple[Optional[float], str]:
    """Return value and a status that distinguishes blank from numeric zero."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing"
    try:
        return float(value), "numeric"
    except (TypeError, ValueError):
        return None, "invalid"


def _week_start(value: date) -> str:
    return (value - timedelta(days=value.weekday())).isoformat()


def _port_name(value: Any) -> str:
    text = _text(value)
    for suffix in ("港区", "港"):
        if text.endswith(suffix) and len(text) > len(suffix):
            return text[: -len(suffix)]
    return text


def _sample_info(raw_sample: Any) -> tuple[str, str, str, str]:
    sample = _text(raw_sample)
    english = ("Jiangyin", "Nantong", "Taicang", "Zhoushan", "Fuzhou", "Zhanjiang", "Lanqiao", "Lanshan", "Lianyungang", "Qingdao", "Rizhao", "Caofeidian", "Huanghua", "Jingtang", "Tianjin")
    names = {name.lower(): port for name, (port, _) in zip(english, SAMPLE_TO_PORT.values())}
    names.update({"连云": "连云港", "天津港": "天津", "total": "总计"})
    normalized = names.get(sample.lower(), sample)
    port, region = SAMPLE_TO_PORT.get(sample, (normalized, ""))
    return sample, port, region, "sample" if sample in SAMPLE_TO_PORT else "unknown_sample"


def _is_total_label(value: Any) -> bool:
    text = _text(value)
    return text.lower() == "total" or text in {"总计", "全国", "全国合计"} or "总计" in text or text.endswith("合计")


def _header_row(rows: Sequence[Sequence[Any]], first_value: str) -> Optional[int]:
    for idx, row in enumerate(rows):
        if row and _text(row[0]) == first_value:
            return idx
    return None


def _read_sheet(path: Path, sheet_name: str) -> List[List[Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return []
        return [list(row) for row in wb[sheet_name].iter_rows(values_only=True)]
    finally:
        wb.close()


def _source_cell(row_no: int, col_no: int) -> str:
    # Excel columns are short in these templates; avoid adding an extra
    # dependency just to format the address.
    value = col_no
    letters = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row_no}"


def _base_row(
    *,
    path: Path,
    sheet: str,
    row_no: int,
    col_no: int,
    observed: Optional[date],
    raw_port: Any,
    value: Any,
    mapping_version: str,
) -> Dict[str, Any]:
    number, status = _number(value)
    sample, port, region, sample_status = _sample_info(raw_port)
    return {
        "observed_date": observed.isoformat() if observed else "",
        "week_start": _week_start(observed) if observed else "",
        "sample_name": sample,
        "port_name": port,
        "region": region,
        "scope_type": "total" if _is_total_label(raw_port) else "sample",
        "value": number,
        "value_status": status,
        "sample_mapping_status": sample_status,
        "source_file": path.name,
        "source_sheet": sheet,
        "source_row": row_no,
        "source_column": col_no,
        "source_cell": _source_cell(row_no, col_no),
        "mapping_version": mapping_version,
        "unit": "万吨",
    }


def _inventory_identity(sheet: str, raw_product: str, category: str = "") -> tuple[str, str, str]:
    # Import lazily: data_visualization already imports database helpers and
    # importing it at module load time would create a circular dependency.
    from . import data_visualization as dv

    source, product, canonical_category = dv._canonical_inventory_identity(
        sheet, raw_product, category, ""
    )
    return source, product, canonical_category


def _category_for_product_sheet(sheet: str) -> str:
    return {"粉矿": "粉矿", "粗粉": "粉矿", "块矿": "块矿", "球团": "球团", "精粉": "精粉"}.get(sheet, "")


def _parse_inventory(path: Path, package: SourcePackage) -> None:
    historical = bool(package.source_files and package.source_files[-1]["template_type"] == "inventory_history")
    excluded_dates = set()
    for sheet_name in ("总览", "粗粉", "粉矿", "块矿", "球团", "精粉", "分品位", "主流品种"):
        rows = _read_sheet(path, sheet_name)
        if not rows:
            continue
        header_index = _header_row(rows, "统计日期")
        if header_index is None:
            continue
        headers = [_text(value) for value in rows[header_index]]
        # Historical workbooks contain a repeated horizontal auxiliary block.
        if "统计日期" in headers[1:]:
            headers = headers[:headers.index("统计日期", 1)]
        old_three_grades = any(h.startswith("中品") for h in headers)
        if historical:
            seen_headers = set()
            for index, header in enumerate(headers):
                if header in seen_headers:
                    headers[index] = ""
                elif header:
                    seen_headers.add(header)
        latest_date = None
        for offset, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if len(values) < 2:
                continue
            observed = _date_value(values[0])
            if observed is None:
                continue
            if historical and sheet_name == "总览":
                if latest_date is not None and observed < latest_date:
                    excluded_dates.add(observed.isoformat())
                    continue
                latest_date = observed
            if observed.isoformat() in excluded_dates:
                continue
            raw_port = values[1]
            if not _text(raw_port):
                continue
            sample, port, region, sample_status = _sample_info(raw_port)

            if sheet_name == "总览":
                for col_index, header in enumerate(headers[3:], start=4):
                    if header not in {"库存总量", "粉矿", "块矿", "球团", "精粉"}:
                        continue
                    row = _base_row(
                        path=path, sheet=sheet_name, row_no=offset, col_no=col_index,
                        observed=observed, raw_port=raw_port,
                        value=values[col_index - 1] if col_index <= len(values) else None,
                        mapping_version=package.mapping_version,
                    )
                    row.update({
                        "metric": header,
                        "sample_name": sample,
                        "port_name": port,
                        "region": region,
                        "sample_mapping_status": sample_status,
                    })
                    package.inventory_summary.append(row)
                continue

            if sheet_name in {"粗粉", "粉矿", "块矿", "球团", "精粉"}:
                category = _category_for_product_sheet(sheet_name)
                for col_index, raw_product in enumerate(headers[3:], start=4):
                    if not raw_product or "总计" in raw_product:
                        continue
                    source_country, product, canonical_category = _inventory_identity(
                        "粗粉" if sheet_name == "粉矿" else sheet_name, raw_product, category
                    )
                    row = _base_row(
                        path=path, sheet=sheet_name, row_no=offset, col_no=col_index,
                        observed=observed, raw_port=raw_port,
                        value=values[col_index - 1] if col_index <= len(values) else None,
                        mapping_version=package.mapping_version,
                    )
                    row.update({
                        "raw_product": raw_product,
                        "source_country": source_country,
                        "product": product,
                        "category": canonical_category,
                        "mainstream_status": _mainstream_status(product, canonical_category),
                        "sample_name": sample,
                        "port_name": port,
                        "region": region,
                        "sample_mapping_status": sample_status,
                    })
                    package.inventory_port_product.append(row)
                continue

            if sheet_name == "分品位":
                for col_index, raw_grade in enumerate(headers[3:], start=4):
                    grade, category = _grade_and_category(raw_grade)
                    if not grade:
                        continue
                    if old_three_grades:
                        grade = {"高品": "高品（旧三档）", "中品": "中低品"}.get(grade, grade)
                    row = _base_row(
                        path=path, sheet=sheet_name, row_no=offset, col_no=col_index,
                        observed=observed, raw_port=raw_port,
                        value=values[col_index - 1] if col_index <= len(values) else None,
                        mapping_version=package.mapping_version,
                    )
                    row.update({
                        "raw_grade": raw_grade,
                        "grade": grade,
                        "category": category,
                        "metric": "库存",
                        "sample_name": sample,
                        "port_name": port,
                        "region": region,
                        "sample_mapping_status": sample_status,
                    })
                    package.inventory_grade.append(row)
                continue

            if sheet_name == "主流品种":
                # This sheet contains ownership sub-columns and a product
                # total.  Only the explicit 汇总 columns are captured here;
                # ownership columns remain available in the source archive.
                for col_index, raw_label in enumerate(headers[3:], start=4):
                    if not raw_label or "汇总" not in raw_label:
                        continue
                    raw_product = raw_label.replace("汇总", "").strip()
                    source_country, product, category = _inventory_identity("粗粉", raw_product, "粉矿")
                    row = _base_row(
                        path=path, sheet=sheet_name, row_no=offset, col_no=col_index,
                        observed=observed, raw_port=raw_port,
                        value=values[col_index - 1] if col_index <= len(values) else None,
                        mapping_version=package.mapping_version,
                    )
                    row.update({
                        "raw_product": raw_product,
                        "source_country": source_country,
                        "product": product,
                        "category": category,
                        "mainstream_status": "主流",
                        "sample_name": sample,
                        "port_name": port,
                        "region": region,
                        "sample_mapping_status": sample_status,
                    })
                    package.inventory_mainstream.append(row)

    if excluded_dates:
        package.validation.setdefault("excluded_out_of_order_dates", []).extend(sorted(excluded_dates))


def _grade_and_category(label: str) -> tuple[str, str]:
    label = _text(label)
    if not label:
        return "", ""
    grade = next((prefix for prefix in ("中高品", "中低品", "高品", "中品", "低品") if label.startswith(prefix)), "")
    if not grade:
        return "", ""
    if label.endswith("总计"):
        return grade, "全品种"
    category = next((item for item in ("粉矿", "块矿", "球团", "精粉") if item in label), "")
    return grade, category


def _arrival_product(raw_product: str) -> tuple[str, str, str]:
    from . import data_visualization as dv

    text = _text(raw_product)
    if text == "总计":
        return "全品种", "全品种", "总计"
    direct = {
        "大杨迪": ("杨迪粉", "粉矿"),
        "PB粉": ("PB粉", "粉矿"),
        "PB块": ("PB块", "块矿"),
        "卡粉": ("卡粉", "粉矿"),
        "巴混": ("巴混", "粉矿"),
        "金布巴粉": ("金布巴粉", "粉矿"),
        "麦克粉": ("麦克粉", "粉矿"),
        "纽曼粉": ("纽曼粉", "粉矿"),
    }
    if text in direct:
        product, category = direct[text]
    elif "块" in text:
        product, category = text, "块矿"
    elif "球" in text:
        product, category = text, "球团"
    elif "精粉" in text:
        product, category = text, "精粉"
    elif any(token in text for token in ("粉", "混", "卡")):
        product, category = text, "粉矿"
    else:
        product, category = text, "未知"
    source_country, canonical, canonical_category = dv._canonical_inventory_identity(
        "粗粉" if category == "粉矿" else category, product, category, ""
    )
    return canonical or product, canonical_category or category, source_country


def _parse_actual_arrival(path: Path, package: SourcePackage) -> None:
    for sheet_name, slice_type in (("国家", "country"), ("品种", "product"), ("货种品位", "form")):
        rows = _read_sheet(path, sheet_name)
        if not rows:
            continue
        header_index = _header_row(rows, "到港时间")
        if header_index is None:
            continue
        headers = [_text(value) for value in rows[header_index]]
        form_by_column = {}
        if slice_type == "form" and header_index > 0:
            parent_form = ""
            parent_headers = rows[header_index - 1]
            for index in range(2, len(headers)):
                parent = _text(parent_headers[index]) if index < len(parent_headers) else ""
                if parent:
                    parent_form = next((form for form in ("粉矿", "精粉", "块矿", "球团", "未知") if form in parent), "")
                form_by_column[index + 1] = parent_form
        for offset, values in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if len(values) < 2:
                continue
            observed = _date_value(values[0])
            if observed is None or not _text(values[1]):
                continue
            raw_port = values[1]
            sample, port, region, sample_status = _sample_info(raw_port)
            if sample_status == "unknown_sample":
                port = _port_name(raw_port)
                region = ""
            scope_type = "total" if _is_total_label(raw_port) else "port"
            for col_index, raw_dimension in enumerate(headers[2:], start=3):
                raw_dimension = _text(raw_dimension)
                if not raw_dimension:
                    continue
                raw_value = values[col_index - 1] if col_index <= len(values) else None
                row = _base_row(
                    path=path, sheet=sheet_name, row_no=offset, col_no=col_index,
                    observed=observed, raw_port=raw_port, value=raw_value,
                    mapping_version=package.mapping_version,
                )
                row.update({
                    "arrival_kind": "actual",
                    "method": "reported_47_port",
                    "slice_type": slice_type,
                    "dimension": raw_dimension,
                    "scope_type": scope_type,
                    "sample_name": sample,
                    "port_name": port,
                    "region": region,
                    "sample_mapping_status": sample_status,
                    "raw_dimension": raw_dimension,
                    "grade": "",
                    "product": "",
                    "category": "",
                    "source_country": "",
                    "mainstream_status": "",
                })
                if slice_type == "country":
                    row["source_country"] = raw_dimension
                elif slice_type == "product":
                    product, category, source_country = _arrival_product(raw_dimension)
                    row.update({
                        "product": product,
                        "category": category,
                        "source_country": source_country,
                        "mainstream_status": _mainstream_status(product, category),
                    })
                else:
                    context = form_by_column.get(col_index, "")
                    grade, category = _arrival_form_dimension(context + raw_dimension)
                    row.update({"grade": grade, "category": category or "未知"})
                package.arrival_actual.append(row)


def _arrival_form_dimension(label: str) -> tuple[str, str]:
    text = _text(label)
    if "汇总" in text:
        if "粉矿" in text:
            return "", "粉矿"
        if "精粉" in text:
            return "", "精粉"
        if "块矿" in text:
            return "", "块矿"
        if "球团" in text:
            return "", "球团"
    if "以下" in text:
        return "低品", "粉矿" if "粉" in text else "块矿" if "块" in text else "未知"
    if "62%以上" in text or "62%以上" in text:
        return "高品", "粉矿" if "粉" in text else "块矿" if "块" in text else "精粉"
    if "60%-62%" in text:
        return "中品", "粉矿" if "粉" in text else "块矿"
    if "球团" in text:
        return "高品", "球团"
    return "", "未知"


def _point_to_row(point: Dict[str, Any], *, kind: str, method: str) -> Dict[str, Any]:
    return {
        "arrival_kind": kind,
        "method": method,
        "observed_date": point.get("display_date", ""),
        "week_start": point.get("week_start", ""),
        "port_name": "",
        "scope_type": "national",
        "slice_type": "product",
        "dimension": point.get("raw_product") or point.get("product", ""),
        "product": point.get("product", ""),
        "category": point.get("category", ""),
        "grade": "",
        "source_country": point.get("source_country", ""),
        "mainstream_status": point.get("mainstream_status", ""),
        "value": point.get("value"),
        "value_status": "numeric" if point.get("value") is not None else "missing",
        "unit": point.get("unit", "万吨"),
        "source_file": point.get("source_file", ""),
        "source_sheet": point.get("source_sheet", ""),
        "source_row": point.get("source_row"),
        "source_column": point.get("source_column"),
        "source_cell": point.get("source_cell", ""),
        "mapping_version": MAPPING_VERSION,
    }


def _parse_estimated_and_shipments(paths: Sequence[Path], package: SourcePackage) -> None:
    from . import data_visualization as dv

    for path in paths:
        if not path.exists():
            continue
        for point in dv._extract_australia_arrivals(path):
            package.arrival_estimated.append(_point_to_row(point, kind="source_forecast", method="australia_anchor_forecast"))
        for point in dv._extract_brazil_estimated_arrivals(path):
            package.arrival_estimated.append(_point_to_row(point, kind="model_estimate", method="brazil_shipments_lag_6w_75pct"))
        for point in dv._extract_australia_shipments(path):
            package.shipments.append(_point_to_row(point, kind="shipment", method="australia_reported_shipment"))
        for point in dv._extract_brazil_card_powder_shipments(path):
            package.shipments.append(_point_to_row(point, kind="shipment", method="brazil_card_powder_reported_shipment"))
        for point in dv._extract_global_shipments(path):
            package.shipments.append(_point_to_row(point, kind="shipment", method="global_reported_shipment"))


def _classify_file(path: Path, sheets: Sequence[str]) -> str:
    names = set(sheets)
    if {"总览", "粉矿"}.issubset(names) and "粗粉" not in names:
        return "inventory_history"
    if {"总览", "粗粉"}.issubset(names):
        return "inventory"
    if {"国家", "品种", "货种品位"}.issubset(names):
        return "arrival_actual"
    if "澳洲预计到达中国锚地量" in names or "巴西发货量" in names:
        return "shipment_and_estimate"
    if "全球铁矿石发运量" in names:
        return "global_shipment"
    return "unknown"


def _file_metadata(path: Path) -> Dict[str, Any]:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = list(wb.sheetnames)
    wb.close()
    return {
        "file_name": path.name,
        "path": str(path),
        "sha256": digest,
        "template_type": _classify_file(path, sheets),
        "template_version": "Mysteel新版" if "新版" in path.name else "Mysteel",
        "sheets": sheets,
    }


def parse_mysteel_source_files(paths: Sequence[Path | str]) -> SourcePackage:
    """Parse original workbooks into a non-activated source package."""
    normalized = [Path(path) for path in paths]
    package = SourcePackage()
    for path in normalized:
        if not path.exists():
            package.validation.setdefault("missing_files", []).append(path.name)
            continue
        package.source_files.append(_file_metadata(path))
        sheets = package.source_files[-1]["sheets"]
        if "总览" in sheets and ({"粗粉", "粉矿"} & set(sheets)):
            _parse_inventory(path, package)
        if {"国家", "品种", "货种品位"}.issubset(set(sheets)):
            _parse_actual_arrival(path, package)
    supply_paths = [Path(item["path"]) for item in package.source_files
                    if item["template_type"] in {"shipment_and_estimate", "global_shipment"}]
    _parse_estimated_and_shipments(supply_paths, package)

    # Keep the V1 summary in the exported workbook so existing charts and
    # table-demand consumers can continue to read the same columns.  This is
    # still only a package preview; it does not touch analytical tables.
    try:
        from . import data_visualization as dv

        # The old historical format is captured by V2 facts. Keep the V1 parser
        # on its supported source formats rather than partially importing history.
        legacy_paths = [Path(item["path"]) for item in package.source_files
                        if item["template_type"] != "inventory_history"]
        package.legacy_points = dv.integrate_mysteel_files(legacy_paths).get("points", [])
    except Exception as exc:  # source detail remains useful even if V1 parse warns
        package.validation.setdefault("legacy_parse_warnings", []).append(str(exc))

    all_dates = [row.get("observed_date") for row in package.all_rows if row.get("observed_date")]
    package.validation.update({
        "source_file_count": len(package.source_files),
        "inventory_port_product_count": len(package.inventory_port_product),
        "inventory_summary_count": len(package.inventory_summary),
        "inventory_grade_count": len(package.inventory_grade),
        "inventory_mainstream_count": len(package.inventory_mainstream),
        "arrival_actual_count": len(package.arrival_actual),
        "arrival_estimated_count": len(package.arrival_estimated),
        "shipment_count": len(package.shipments),
        "date_min": min(all_dates) if all_dates else "",
        "date_max": max(all_dates) if all_dates else "",
        "missing_value_count": sum(1 for row in package.all_rows if row.get("value_status") == "missing"),
        "invalid_value_count": sum(1 for row in package.all_rows if row.get("value_status") == "invalid"),
    })
    return package


def archive_source_package(package: SourcePackage, user_name: str) -> Dict[str, Any]:
    """Copy original files and record hashes before temporary uploads are removed."""
    archive_root = db.DATA_DIR / "iron_ore_source_archive" / package.package_id
    archive_root.mkdir(parents=True, exist_ok=True)
    source_ids: List[int] = []
    with db.connect() as conn:
        cur = conn.cursor()
        for metadata in package.source_files:
            source_path = Path(metadata["path"])
            destination = archive_root / metadata["file_name"]
            if source_path.resolve() != destination.resolve():
                shutil.copy2(source_path, destination)
            metadata["archive_path"] = str(destination)
            metadata["path"] = str(destination)
            existing = db._exec(cur, "SELECT id FROM dv_source_files WHERE sha256 = ?", (metadata["sha256"],)).fetchone()
            if existing:
                source_ids.append(existing["id"])
                continue
            source_id = db._last_insert_id(
                cur,
                """INSERT INTO dv_source_files
                   (file_name, sha256, archive_path, template_type, template_version,
                    source_date_start, source_date_end, uploaded_by, metadata_json)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    metadata["file_name"], metadata["sha256"], str(destination),
                    metadata["template_type"], metadata.get("template_version"),
                    package.validation.get("date_min", ""), package.validation.get("date_max", ""),
                    user_name, json.dumps(metadata, ensure_ascii=False),
                ),
            )
            source_ids.append(source_id)
        db._exec(
            cur,
            """INSERT INTO dv_source_packages
               (package_id, structure_version, parser_version, mapping_version, status,
                source_file_ids, validation_json, created_by)
               VALUES (?, ?, ?, ?, 'prepared', ?, ?, ?)""",
            (
                package.package_id,
                package.structure_version,
                package.parser_version,
                package.mapping_version,
                json.dumps(source_ids, ensure_ascii=False),
                json.dumps(package.validation, ensure_ascii=False),
                user_name,
            ),
        )
        conn.commit()
    return {
        "package_id": package.package_id,
        "archive_root": str(archive_root),
        "source_file_ids": source_ids,
        "source_files": package.source_files,
    }


def update_source_package_output(package_id: str, output_path: str, output_sha256: str) -> None:
    with db.connect() as conn:
        cur = conn.cursor()
        db._exec(
            cur,
            "UPDATE dv_source_packages SET output_path = ?, output_sha256 = ? WHERE package_id = ?",
            (output_path, output_sha256, package_id),
        )
        conn.commit()


def _insert_fact_rows(cur, table: str, columns: Sequence[str], rows: Sequence[Dict[str, Any]]) -> int:
    if not rows:
        return 0
    identity_columns = [column for column in (
        "observed_date", "week_start", "port_name", "scope_type", "product", "category",
        "source_country", "metric", "grade", "raw_grade", "arrival_kind", "method", "slice_type", "dimension",
    ) if column in columns]
    if db._is_pg():
        # Serialize overlapping imports while comparing and inserting facts.
        db._exec(cur, f"LOCK TABLE {table} IN SHARE ROW EXCLUSIVE MODE")
    dates = sorted({row.get("observed_date") for row in rows if row.get("observed_date")})
    existing = []
    for start in range(0, len(dates), 500):
        chunk = dates[start:start + 500]
        placeholders = ", ".join("?" for _ in chunk)
        existing.extend(dict(row) for row in db._exec(
            cur, f"SELECT * FROM {table} WHERE observed_date IN ({placeholders})", tuple(chunk)
        ).fetchall())

    def identity(row):
        return tuple(row.get(column) or "" for column in identity_columns)

    known = {identity(row): row for row in existing}
    unique = []
    for row in rows:
        key = identity(row)
        prior = known.get(key)
        if prior is not None:
            if (prior.get("value") != row.get("value")
                    or prior.get("value_status") != row.get("value_status")
                    or prior.get("unit") != row.get("unit")):
                raise ValueError(
                    f"同一观测记录存在数值冲突：{row.get('observed_date')} "
                    f"{row.get('port_name')} {row.get('product') or row.get('dimension') or row.get('metric') or row.get('grade')}；"
                    "本次导入已取消，请核对来源版本"
                )
            continue
        known[key] = row
        unique.append(row)
    placeholders = ", ".join("?" for _ in columns)
    sql = f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({placeholders})"
    db._executemany(cur, sql, [tuple(row.get(column) for column in columns) for row in unique])
    return len(unique)


def _store_source_package_facts_in_connection(
    cur,
    package: SourcePackage,
    source_file_ids: Sequence[int] = (),
    user_name: str = "",
) -> Dict[str, Any]:
    """Write one package's detail facts using the caller's transaction."""
    existing = db._exec(
        cur, "SELECT id, status FROM dv_source_packages WHERE package_id = ?", (package.package_id,)
    ).fetchone()
    if existing:
        if existing["status"] == "activated":
            return {"duplicate": True, "package_id": package.package_id, "inserted": 0}
    else:
        db._exec(
            cur,
            """INSERT INTO dv_source_packages
               (package_id, structure_version, parser_version, mapping_version, status,
                source_file_ids, validation_json, created_by)
               VALUES (?, ?, ?, ?, 'prepared', ?, ?, ?)""",
            (
                package.package_id,
                package.structure_version,
                package.parser_version,
                package.mapping_version,
                json.dumps(list(source_file_ids), ensure_ascii=False),
                json.dumps(package.validation, ensure_ascii=False),
                user_name,
            ),
        )

    common_inventory = (
        "package_id", "observed_date", "week_start", "sample_name", "port_name", "region", "scope_type",
        "raw_product", "product", "category", "source_country", "mainstream_status", "value", "value_status",
        "unit", "source_file", "source_sheet", "source_row", "source_column", "source_cell", "mapping_version",
    )
    inventory_rows = [{**row, "package_id": package.package_id} for row in package.inventory_port_product]
    inventory_count = _insert_fact_rows(cur, "dv_port_inventory_facts", common_inventory, inventory_rows)
    mainstream_count = _insert_fact_rows(
        cur,
        "dv_inventory_mainstream_facts",
        (
            "package_id", "observed_date", "week_start", "sample_name", "port_name", "region", "scope_type",
            "raw_product", "product", "category", "source_country", "value", "value_status", "unit",
            "source_file", "source_sheet", "source_row", "source_column", "source_cell", "mapping_version",
        ),
        [{**row, "package_id": package.package_id} for row in package.inventory_mainstream],
    )
    summary_count = _insert_fact_rows(
        cur,
        "dv_inventory_summary_facts",
        (
            "package_id", "observed_date", "week_start", "sample_name", "port_name", "region", "scope_type",
            "metric", "value", "value_status", "unit", "source_file", "source_sheet", "source_row",
            "source_column", "source_cell", "mapping_version",
        ),
        [{**row, "package_id": package.package_id} for row in package.inventory_summary],
    )
    grade_count = _insert_fact_rows(
        cur,
        "dv_inventory_grade_facts",
        (
            "package_id", "observed_date", "week_start", "sample_name", "port_name", "region", "scope_type",
            "raw_grade", "grade", "category", "value", "value_status", "unit", "source_file", "source_sheet",
            "source_row", "source_column", "source_cell", "mapping_version",
        ),
        [{**row, "package_id": package.package_id} for row in package.inventory_grade],
    )
    arrival_rows = [{**row, "package_id": package.package_id} for row in [*package.arrival_actual, *package.arrival_estimated]]
    arrival_count = _insert_fact_rows(
        cur,
        "dv_arrival_facts",
        (
            "package_id", "arrival_kind", "method", "observed_date", "week_start", "port_name", "scope_type",
            "slice_type", "dimension", "product", "category", "grade", "source_country", "mainstream_status",
            "value", "value_status", "unit", "source_file", "source_sheet", "source_row", "source_column",
            "source_cell", "mapping_version",
        ),
        arrival_rows,
    )
    db._exec(
        cur,
        "UPDATE dv_source_packages SET status = 'activated' WHERE package_id = ?",
        (package.package_id,),
    )
    return {
        "duplicate": False,
        "package_id": package.package_id,
        "inserted": inventory_count + mainstream_count + summary_count + grade_count + arrival_count,
        "inventory_port_product": inventory_count,
        "inventory_mainstream": mainstream_count,
        "inventory_summary": summary_count,
        "inventory_grade": grade_count,
        "arrival": arrival_count,
    }


def store_source_package_facts(
    package: SourcePackage,
    source_file_ids: Sequence[int] = (),
    user_name: str = "",
) -> Dict[str, Any]:
    """Activate parsed detail facts for a package, once."""
    with db.connect() as conn:
        result = _store_source_package_facts_in_connection(
            conn.cursor(), package, source_file_ids, user_name
        )
        conn.commit()
    return result


def _append_rows(sheet, headers: Sequence[str], rows: Iterable[Dict[str, Any]]) -> None:
    sheet.append(list(headers))
    for row in rows:
        sheet.append([row.get(header) for header in headers])


def build_v2_workbook(package: SourcePackage) -> bytes:
    """Build the downloadable, versioned integrated workbook."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook(write_only=True)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)

    def styled_sheet(name: str):
        sheet = wb.create_sheet(name)
        # WriteOnlyCell keeps the output small for a full historical package.
        from openpyxl.cell import WriteOnlyCell

        cells = []
        for label in headers_by_sheet[name]:
            cell = WriteOnlyCell(sheet, value=label)
            cell.font = font
            cell.fill = fill
            cells.append(cell)
        sheet.append(cells)
        return sheet

    headers_by_sheet = {
        "整合明细": [
            "统计周一", "统计周日", "业务年份", "业务周次", "周次标签", "展示日期", "数据类型",
            "来源/国家", "品种", "种类", "主流/非主流", "数值", "单位", "来源文件", "来源Sheet",
            "来源区域", "是否参与表需", "校验状态", "备注",
        ],
        "港口品种库存": [
            "观测日期", "统计周一", "样本", "港口", "区域", "范围", "原始品种", "品种",
            "种类", "来源国家", "主流/非主流", "数值", "数值状态", "单位", "来源文件", "来源Sheet",
            "来源行", "来源列", "来源单元格", "映射版本",
        ],
        "主流品种库存": [
            "观测日期", "统计周一", "样本", "港口", "区域", "范围", "原始品种", "品种",
            "种类", "来源国家", "数值", "数值状态", "单位", "来源文件", "来源Sheet", "来源行",
            "来源列", "来源单元格", "映射版本",
        ],
        "库存汇总分档": [
            "观测日期", "统计周一", "样本", "港口", "区域", "范围", "指标", "品位档",
            "种类", "数值", "数值状态", "单位", "来源文件", "来源Sheet", "来源行", "来源列",
            "来源单元格", "映射版本",
        ],
        "到港明细": [
            "到港口径", "方法", "观测日期", "统计周一", "港口", "范围", "切片", "维度",
            "品种", "种类", "品位档", "来源国家", "主流/非主流", "数值", "数值状态", "单位",
            "来源文件", "来源Sheet", "来源行", "来源列", "来源单元格", "映射版本",
        ],
        "批次信息": ["字段", "值"],
    }

    # Legacy summary rows are already normalized by data_visualization.
    legacy = styled_sheet("整合明细")
    for point in package.legacy_points:
        legacy.append([
            point.get("week_start"), point.get("week_end"), point.get("business_year"), point.get("business_week"),
            point.get("week_label"), point.get("display_date"), point.get("metric_type"), point.get("source_country"),
            point.get("product"), point.get("category"), point.get("mainstream_status"), point.get("value"),
            point.get("unit", "万吨"), point.get("source_file"), point.get("source_sheet"), point.get("source_section"),
            "是" if point.get("is_calculable") else "否", point.get("validation_status", "ok"), point.get("note", ""),
        ])

    inv = styled_sheet("港口品种库存")
    for row in package.inventory_port_product:
        inv.append([
            row.get("observed_date"), row.get("week_start"), row.get("sample_name"), row.get("port_name"),
            row.get("region"), row.get("scope_type"), row.get("raw_product"), row.get("product"), row.get("category"),
            row.get("source_country"), row.get("mainstream_status"), row.get("value"), row.get("value_status"),
            row.get("unit"), row.get("source_file"), row.get("source_sheet"), row.get("source_row"),
            row.get("source_column"), row.get("source_cell"), row.get("mapping_version"),
        ])

    mainstream = styled_sheet("主流品种库存")
    for row in package.inventory_mainstream:
        mainstream.append([
            row.get("observed_date"), row.get("week_start"), row.get("sample_name"), row.get("port_name"),
            row.get("region"), row.get("scope_type"), row.get("raw_product"), row.get("product"), row.get("category"),
            row.get("source_country"), row.get("value"), row.get("value_status"), row.get("unit"), row.get("source_file"),
            row.get("source_sheet"), row.get("source_row"), row.get("source_column"), row.get("source_cell"),
            row.get("mapping_version"),
        ])

    grade = styled_sheet("库存汇总分档")
    for row in [*package.inventory_summary, *package.inventory_grade]:
        grade.append([
            row.get("observed_date"), row.get("week_start"), row.get("sample_name"), row.get("port_name"),
            row.get("region"), row.get("scope_type"), row.get("metric", "库存"), row.get("grade", ""),
            row.get("category", ""), row.get("value"), row.get("value_status"), row.get("unit"), row.get("source_file"),
            row.get("source_sheet"), row.get("source_row"), row.get("source_column"), row.get("source_cell"),
            row.get("mapping_version"),
        ])

    arrival = styled_sheet("到港明细")
    for row in [*package.arrival_actual, *package.arrival_estimated]:
        arrival.append([
            row.get("arrival_kind"), row.get("method"), row.get("observed_date"), row.get("week_start"), row.get("port_name"),
            row.get("scope_type"), row.get("slice_type"), row.get("dimension"), row.get("product"), row.get("category"),
            row.get("grade", ""), row.get("source_country", ""), row.get("mainstream_status", ""), row.get("value"),
            row.get("value_status"), row.get("unit"), row.get("source_file"), row.get("source_sheet"), row.get("source_row"),
            row.get("source_column"), row.get("source_cell"), row.get("mapping_version"),
        ])

    info = styled_sheet("批次信息")
    metadata_rows = [
        ("结构版本", package.structure_version),
        ("包ID", package.package_id),
        ("解析版本", package.parser_version),
        ("映射版本", package.mapping_version),
        ("源文件数量", len(package.source_files)),
        ("源文件哈希", json.dumps({item["file_name"]: item["sha256"] for item in package.source_files}, ensure_ascii=False)),
        ("校验结果", json.dumps(package.validation, ensure_ascii=False)),
    ]
    for key, value in metadata_rows:
        info.append([key, value])

    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def _mainstream_status(product: str, category: str) -> str:
    from . import data_visualization as dv

    return dv._mainstream_status(product, category)
